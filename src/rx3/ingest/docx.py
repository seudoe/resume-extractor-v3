"""DOCX -> Document IR (PROMPT.md §5 Stage 5).

DOCX has no page geometry, so bbox is a pseudo-bbox: sequential y per line,
x from paragraph indentation, as PROMPT.md §5 specifies. Hyperlinks are
collected at the document level (from `part.rels`) rather than per-run —
python-docx doesn't expose per-run hyperlink targets without parsing raw
XML, and the header/contact rules (Stage 8) that need them already search
all page links, not a specific line's.
"""

import io

from docx import Document as DocxDocument
from docx.shared import Pt

from ir import BBox, Document, DocumentMeta, Line, Link, Page, Span

LINE_HEIGHT = 14.0  # pt, a reasonable single-spaced default
PAGE_WIDTH = 612.0  # Letter, pt — DOCX has no real page geometry to read
PAGE_HEIGHT = 792.0
DEFAULT_FONT_SIZE = 11.0
BULLET_GLYPHS = ("•", "●", "○", "▪", "–", "-", "*", "➢", "✓")


def _is_heading(style_name: str) -> bool:
    return style_name.lower().startswith("heading") or style_name.lower() == "title"


def _is_list_style(style_name: str) -> bool:
    return "list" in style_name.lower() or "bullet" in style_name.lower()


def _paragraph_text_and_style(paragraph) -> tuple[str, bool, float, str]:
    """Returns (text, bold, size_pt, font_name) using the first non-empty
    run as representative — multi-run paragraphs with mixed styling are a
    simplification accepted here; Stage 7+ works off the merged line text."""
    runs = [r for r in paragraph.runs if r.text.strip()]
    text = "".join(r.text for r in paragraph.runs).strip()
    style_name = paragraph.style.name if paragraph.style else ""
    heading = _is_heading(style_name)

    if runs:
        rep = runs[0]
        bold = bool(rep.bold) or heading
        size = rep.font.size.pt if isinstance(rep.font.size, Pt) else (rep.font.size or DEFAULT_FONT_SIZE)
        font = rep.font.name or ""
    else:
        bold = heading
        size = DEFAULT_FONT_SIZE
        font = ""

    if _is_list_style(style_name) and text and not text.startswith(BULLET_GLYPHS):
        text = f"• {text}"

    return text, bold, float(size), font


def ingest_docx(data: bytes, filename: str = "") -> Document:
    doc = DocxDocument(io.BytesIO(data))

    lines: list[Line] = []
    y = 0.0
    line_counter = 0

    def add_line(text: str, bold: bool, size: float, font: str, indent_pt: float) -> None:
        nonlocal y, line_counter
        if not text.strip():
            return
        x0 = indent_pt
        bbox = BBox(x0=x0, y0=y, x1=x0 + min(len(text) * size * 0.5, PAGE_WIDTH - x0), y1=y + LINE_HEIGHT)
        lines.append(
            Line(
                id=f"L{line_counter}",
                page=0,
                bbox=bbox,
                spans=[Span(text=text, bbox=bbox, size=size, bold=bold, italic=False, font=font)],
                text=text,
                source="docx",
            )
        )
        line_counter += 1
        y += LINE_HEIGHT

    for paragraph in doc.paragraphs:
        text, bold, size, font = _paragraph_text_and_style(paragraph)
        indent = paragraph.paragraph_format.left_indent
        indent_pt = indent.pt if indent else 0.0
        add_line(text, bold, size, font, indent_pt)

    for table in doc.tables:
        for row in table.rows:
            cell_texts = [" ".join(p.text for p in cell.paragraphs).strip() for cell in row.cells]
            row_text = "\t".join(t for t in cell_texts if t)
            add_line(row_text, False, DEFAULT_FONT_SIZE, "", 0.0)

    links = [
        Link(bbox=BBox(x0=0, y0=0, x1=0, y1=0), uri=rel.target_ref)
        for rel in doc.part.rels.values()
        if "hyperlink" in rel.reltype
    ]

    page = Page(number=0, width=PAGE_WIDTH, height=max(y, PAGE_HEIGHT), lines=lines, links=links, drawings=[])
    return Document(pages=[page], meta=DocumentMeta(filename=filename, page_count=1, source="docx"))

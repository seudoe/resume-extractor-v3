"""Minimal PDF/DOCX -> plain text, for gold-drafting only.

This is NOT the Stage 5 ingester (that keeps layout: font size, bold,
bbox). This just dumps readable text so a human (or the agent, drafting
gold offline per PROMPT.md §4.1) has something to read while writing gold
JSON by hand.
"""

from pathlib import Path

import pymupdf as fitz


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path)
    if suffix == ".docx":
        return _extract_docx(path)
    raise ValueError(f"unsupported file type: {suffix}")


def _extract_pdf(path: Path) -> str:
    doc = fitz.open(path)
    try:
        pages = [page.get_text("text") for page in doc]
    finally:
        doc.close()
    return "\n\n--- PAGE BREAK ---\n\n".join(pages)


def extract_links(path: Path) -> list[str]:
    """Pull hyperlink URIs (PROMPT.md §8: links carry the real URL even when
    visible text is an icon). Helps when hand-drafting gold for resumes whose
    link icons get mangled by plain-text extraction."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        doc = fitz.open(path)
        try:
            uris = []
            for page in doc:
                for link in page.get_links():
                    if link.get("uri"):
                        uris.append(link["uri"])
            return uris
        finally:
            doc.close()
    if suffix == ".docx":
        from docx import Document

        doc = Document(path)
        rels = doc.part.rels
        return [r.target_ref for r in rels.values() if "hyperlink" in r.reltype]
    return []


def _extract_docx(path: Path) -> str:
    from docx import Document

    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)

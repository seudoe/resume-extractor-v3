"""Stage 5 exit check: IR snapshot tests for real gold PDFs + DOCX, and the
OCR-flag test.

Real resumes live in ../resume-data/ (sibling dir, not in this git repo,
real PII — see PROMPT.md §1.6/§1.7). Tests that need them skip cleanly when
that directory isn't present, e.g. a fresh clone without the user's data.
"""

import io
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from rx3.ingest.docx import ingest_docx  # noqa: E402
from rx3.ingest.pdf import ingest_pdf  # noqa: E402
from rx3.ingest.quality import page_quality  # noqa: E402

GOLD_PDF_DIR = Path(__file__).resolve().parent.parent.parent / "resume-data" / "PDFs" / "AAA"

# 5 diverse real resumes (PROMPT.md §5 Stage 5 exit check): two-column
# (AltaCV), dense icon-heavy single-column (Simple_Hipster_CV), LaTeX-style
# (Entry_Level...), a real student resume (Sambhav), a name-with-apostrophe
# edge case (Azlan's).
SAMPLE_PDFS = [
    "AltaCV_Template.pdf",
    "Simple_Hipster_CV.pdf",
    "Entry_Level_Resume_Template__LaTeX_.pdf",
    "SambhavMirajgaonkarResume.pdf",
    "Azlan's_Resume-mar26.pdf",
]

requires_gold = pytest.mark.skipif(not GOLD_PDF_DIR.exists(), reason="resume-data/ not present (external, PII)")


@requires_gold
@pytest.mark.parametrize("filename", SAMPLE_PDFS)
def test_pdf_ingest_snapshot(filename):
    path = GOLD_PDF_DIR / filename
    if not path.exists():
        pytest.skip(f"{filename} not found in {GOLD_PDF_DIR}")
    doc = ingest_pdf(path.read_bytes(), filename=filename)

    assert doc.meta.page_count == len(doc.pages) > 0
    all_lines = [line for page in doc.pages for line in page.lines]
    assert len(all_lines) > 5, "expected a real resume to have more than 5 lines"
    # Pointer principle: ids are unique and stable in emission order.
    assert [line.id for line in all_lines] == [f"L{i}" for i in range(len(all_lines))]
    # Every line has at least one span and non-empty text.
    assert all(line.spans and line.text.strip() for line in all_lines)
    # At least one bold line exists on a real resume (name/headings).
    assert any(span.bold for line in all_lines for span in line.spans)


@requires_gold
def test_pdf_ingest_extracts_link_annotations():
    path = GOLD_PDF_DIR / "SambhavMirajgaonkarResume.pdf"
    if not path.exists():
        pytest.skip("SambhavMirajgaonkarResume.pdf not found")
    doc = ingest_pdf(path.read_bytes())
    uris = [link.uri for page in doc.pages for link in page.links]
    assert "https://github.com/sam-wlh-ds" in uris
    assert "https://www.linkedin.com/in/sambhav-m/" in uris


def test_docx_ingest_paragraphs_headings_bullets_tables():
    from docx import Document as DocxDocument

    d = DocxDocument()
    d.add_heading("Alex Johnson", level=1)
    d.add_paragraph("Software Engineer")
    d.add_paragraph("Built things", style="List Bullet")
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skill"
    table.rows[0].cells[1].text = "Python"
    buf = io.BytesIO()
    d.save(buf)

    doc = ingest_docx(buf.getvalue(), filename="sample.docx")
    texts = [line.text for line in doc.pages[0].lines]

    assert "Alex Johnson" in texts
    assert doc.pages[0].lines[0].spans[0].bold is True  # Heading 1 -> bold
    assert any(t.startswith("•") for t in texts)  # bullet glyph added
    assert "Skill	Python" in texts  # table cells tab-separated


def test_docx_ingest_collects_hyperlinks():
    from docx import Document as DocxDocument

    d = DocxDocument()
    d.add_paragraph("placeholder")
    part = d.part
    r_id = part.relate_to(
        "https://github.com/example", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    assert r_id  # relationship created
    buf = io.BytesIO()
    d.save(buf)

    doc = ingest_docx(buf.getvalue())
    assert "https://github.com/example" in [link.uri for link in doc.pages[0].links]


def test_page_quality_flags_near_empty_page_for_ocr():
    """Stand-in for the FlowCV_Resume_2026-08-02.pdf regression test (h.md's
    0-extractable-chars example) — that fixture is currently missing from
    resume-data/ (see PROGRESS.md), so this synthesizes the same failure
    mode: a page whose content is drawn as shapes, not text."""
    import pymupdf

    blank = pymupdf.open()
    page = blank.new_page()
    page.draw_rect(pymupdf.Rect(50, 50, 200, 80))  # vector shape, no text
    data = blank.tobytes()
    blank.close()

    doc = ingest_pdf(data, filename="blank.pdf")
    quality = page_quality(doc.pages[0])
    assert quality["needs_ocr"] is True
    assert quality["char_count"] == 0


@requires_gold
def test_page_quality_does_not_flag_real_text_resume():
    path = GOLD_PDF_DIR / "SambhavMirajgaonkarResume.pdf"
    if not path.exists():
        pytest.skip("SambhavMirajgaonkarResume.pdf not found")
    doc = ingest_pdf(path.read_bytes())
    assert all(page_quality(p)["needs_ocr"] is False for p in doc.pages)


def _image_only_pdf(path):
    import pymupdf

    src = pymupdf.open(path)
    pix = src[0].get_pixmap(dpi=150)
    img_pdf = pymupdf.open()
    page = img_pdf.new_page(width=src[0].rect.width, height=src[0].rect.height)
    page.insert_image(page.rect, pixmap=pix)
    return img_pdf.tobytes()


@requires_gold
@pytest.mark.parametrize("engine", ["tesseract", "rapid"])
def test_ocr_fallback_recovers_text_from_image_only_pdf(engine):
    """Stage 6 exit check: an image-only PDF (no text layer) goes through OCR
    and yields readable text. RapidOCR is slow (~20 s); Tesseract ~2 s."""
    from rx3.ocr import tesseract
    from rx3.ocr.select import apply_ocr_fallback

    if engine == "tesseract" and not tesseract.available():
        pytest.skip("Tesseract not installed")
    path = GOLD_PDF_DIR / "Vedant Patil.pdf"
    if not path.exists():
        pytest.skip("Vedant Patil.pdf not found")
    data = _image_only_pdf(path)

    doc = ingest_pdf(data)
    assert page_quality(doc.pages[0])["needs_ocr"] is True  # no text layer

    doc = apply_ocr_fallback(doc, data, engine=engine)
    text = " ".join(line.text for line in doc.pages[0].lines).upper()
    assert all(line.source == "ocr" for line in doc.pages[0].lines)
    assert "VEDANT" in text and "GMAIL.COM" in text
    assert [l.id for l in doc.pages[0].lines] == [f"L{i}" for i in range(len(doc.pages[0].lines))]
